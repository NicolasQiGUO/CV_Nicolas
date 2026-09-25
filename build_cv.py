"""Génère le CV (DOCX + PDF) à partir du contenu défini dans CV ci-dessous.

Usage : python3 build_cv.py [nom_de_sortie]
Pour une version ciblée : dupliquer ce fichier (ou le dict CV) et n'ajuster que
TITLE, PROFIL, les bullets BNP / La Poste et COMPETENCES.
"""
import copy
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
FONT = "Arial"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x55, 0x55)

CV = {
    "nom": "NICOLAS QI GUO",
    "title": "CHEF DE PROJET SI | TRANSFORMATION DIGITALE, IA & AUTOMATISATION",
    "contact": [
        "Orléans, France | Mobilité Île-de-France",
        "Tél. : +33 6 30 23 24 41 | E-mail : qi.guo@essec.edu",
    ],
    "profil": (
        "Chef de projet SI, major de promotion ESSEC & Télécom Paris et Université Paris Cité, "
        "avec 3 ans d'expérience en transformation digitale chez BNP Paribas et La Poste Groupe. "
        "Pilotage de projets IA (GenAI, RAG, Copilot) et d'automatisation (RPA, Power Platform), "
        "du recueil des besoins à la mise en production, gouvernance projet et conduite du changement."
    ),
    "formation": [
        ("2024 – 2025",
         "Mastère spécialisé en Management des Systèmes d'Information en Réseaux",
         " – Major de promotion", "ESSEC Business School & Télécom Paris"),
        ("2022 – 2023",
         "Master 2 en Projets Informatiques et Stratégie d'entreprise",
         " – Major de promotion", "Université Paris Cité"),
        ("Avant 2016",
         "Master en Littératures Françaises",
         " – Université Paris-Sorbonne", None),
        ("",
         "Licence en Langue Française",
         " – Xi'an JiaoTong University, Chine", None),
    ],
    "experiences": [
        {
            "dates": "2024 – 2025",
            "poste": "Digital Transformation Project Manager – AIR TECH",
            "entreprise": "BNP PARIBAS",
            "contrat": "(Alternance)",
            "perimetre": "Périmètre : Compliance & Risk (AML, sanctions, contrôle et monitoring, sécurité financière)",
            "bullets": [
                ("Structuration et coordination des initiatives IA & RPA : ",
                 "co-construction d'un dispositif de qualification et de priorisation des use cases "
                 "(valeur business, ROI, faisabilité technique) ; 20+ use cases identifiés, "
                 "dont 3 retenus pour un lancement en projet", []),
                ("Pilotage et gouvernance des projets IA : ",
                 "recueil et formalisation des besoins, coordination transverse métiers / IT, "
                 "conduite de POC, coordination et suivi de l'industrialisation de 3 solutions IA "
                 "(GenAI, RAG, Copilot), reporting auprès des instances de gouvernance", []),
                ("Transformation digitale et conduite du changement : ",
                 "contribution au déploiement d'un programme IA à grande échelle (5 000+ collaborateurs), "
                 "animation d'un réseau de 70+ ambassadeurs IA, sessions d'onboarding et d'acculturation, "
                 "événements communautaires", []),
            ],
        },
        {
            "dates": "2022 – 2024",
            "poste": "Chef de Projet SI – Data Intelligence & Innovation DTSFG",
            "entreprise": "LA POSTE GROUPE",
            "contrat": "(Alternance + CDD)",
            "perimetre": "Périmètre : Direction de la Transformation des Solutions Finance (SI Finance, SIRH, SAP)",
            "bullets": [
                ("Pilotage de projets d'automatisation (RPA & Power Platform), dont 2 de bout en bout :", "", [
                    "Recueil et analyse des besoins métiers, rédaction des cahiers des charges et spécifications fonctionnelles",
                    "Coordination métiers / IT pour la coconception des solutions",
                    "Pilotage du développement, de la recette et de la mise en production",
                    "Gestion des risques et animation des instances projets (COSUI, COPIL)",
                    "Accompagnement du déploiement auprès des fonctions Finance et RH",
                ]),
                ("Phase RUN :", "", [
                    "Monitoring quotidien d'une dizaine de robots en production et gestion des incidents",
                    "Suivi de la performance des solutions déployées (ROI, fiabilité, adoption)",
                    "Identification et mise en œuvre d'axes d'amélioration continue",
                ]),
            ],
        },
        {
            "dates": "2018 – 2022",
            "poste": "Retail & VIP Guest Relations",
            "entreprise": "LOUIS VUITTON",
            "contrat": "",
            "perimetre": None,
            "bullets": [
                ("", "Gestion d'une clientèle VIP internationale dans un environnement exigeant", []),
                ("", "Sens du service, communication et adaptation en contexte multiculturel", []),
            ],
        },
    ],
    "competences": [
        ("Gestion de projet : ",
         "cadrage, recueil des besoins, cahier des charges, spécifications fonctionnelles, planification, "
         "coordination métiers / IT, recette, mise en production, gestion des risques, COPIL / COSUI"),
        ("Méthodologies : ", "Cycle en V, Agile, hybride"),
        ("Transformation & IA : ",
         "GenAI, RAG, Copilot, RPA, POC, qualification de use cases, ROI, conduite du changement"),
        ("Environnement SI : ",
         "SI Finance, SIRH, SAP, intégration applicative, sécurité, conformité et gouvernance des SI"),
        ("Outils : ", "Power Apps, Power Automate, Power BI, SharePoint"),
        ("Langages (notions) : ", "Python, SQL, Java, PHP, C, C#, HTML/CSS"),
    ],
    "langues": "Anglais : professionnel (TOEIC 955/990) | Français : bilingue | Chinois : langue maternelle",
    "interets": "Sport (crossfit, natation), voyages (18 pays visités), engagement associatif pour la protection animale",
}


def set_spacing(p, before=0, after=0, line=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def run(p, text, size=9.5, bold=False, italic=False, color=None):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3A5F")):
        b.set(qn(k), v)
    bdr.append(b)
    pPr.append(bdr)


def section(doc, title):
    p = doc.add_paragraph()
    set_spacing(p, before=7, after=3)
    run(p, title, size=10.5, bold=True, color=ACCENT)
    bottom_border(p)


def bullet(doc, label, text, level=0):
    p = doc.add_paragraph()
    set_spacing(p, after=1, line=1.05)
    pf = p.paragraph_format
    indent = 0.5 + 0.5 * level
    pf.left_indent = Cm(indent)
    pf.first_line_indent = Cm(-0.35)
    run(p, ("• " if level == 0 else "– ") + " ")
    if label:
        run(p, label, bold=True)
    if text:
        run(p, text)


def float_picture(paragraph, path, width_cm, right_offset_cm=0, top_offset_cm=0):
    """Ajoute une image flottante (ancrée à droite) sans casser le flux texte."""
    r = paragraph.add_run()
    inline_shape = r.add_picture(str(path), width=Cm(width_cm))
    inline = inline_shape._inline
    cx, cy = inline.extent.cx, inline.extent.cy
    anchor = OxmlElement("wp:anchor")
    for k, v in (("distT", "0"), ("distB", "0"), ("distL", "114300"), ("distR", "0"),
                 ("simplePos", "0"), ("relativeHeight", "251658240"), ("behindDoc", "0"),
                 ("locked", "0"), ("layoutInCell", "1"), ("allowOverlap", "1")):
        anchor.set(k, v)
    sp = OxmlElement("wp:simplePos"); sp.set("x", "0"); sp.set("y", "0"); anchor.append(sp)
    ph = OxmlElement("wp:positionH"); ph.set("relativeFrom", "margin")
    al = OxmlElement("wp:align"); al.text = "right"; ph.append(al); anchor.append(ph)
    pv = OxmlElement("wp:positionV"); pv.set("relativeFrom", "paragraph")
    po = OxmlElement("wp:posOffset"); po.text = str(int(Cm(top_offset_cm))); pv.append(po); anchor.append(pv)
    ext = OxmlElement("wp:extent"); ext.set("cx", str(cx)); ext.set("cy", str(cy)); anchor.append(ext)
    ee = OxmlElement("wp:effectExtent")
    for k in ("l", "t", "r", "b"):
        ee.set(k, "0")
    anchor.append(ee)
    wrap = OxmlElement("wp:wrapSquare"); wrap.set("wrapText", "bothSides"); anchor.append(wrap)
    for child in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        el = inline.find(qn(child))
        if el is not None:
            anchor.append(copy.deepcopy(el))
    inline.getparent().replace(inline, anchor)


def build(out_stem):
    doc = Document()
    s = doc.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.top_margin = s.bottom_margin = Cm(1.2)
    s.left_margin = s.right_margin = Cm(1.4)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9.5)

    # En-tête : nom, title, contact + photo flottante
    p = doc.add_paragraph()
    set_spacing(p, after=2)
    photo = ROOT / "assets" / "photo.png"
    if photo.exists():
        float_picture(p, photo, width_cm=2.9)
    run(p, CV["nom"], size=18, bold=True, color=ACCENT)

    p = doc.add_paragraph(); set_spacing(p, after=4)
    run(p, CV["title"], size=11, bold=True)
    for line in CV["contact"]:
        p = doc.add_paragraph(); set_spacing(p, after=0)
        run(p, line, size=9.5, color=GREY)

    p = doc.add_paragraph(); set_spacing(p, before=6, after=0, line=1.1)
    run(p, "Profil : ", bold=True, color=ACCENT)
    run(p, CV["profil"])

    section(doc, "EXPÉRIENCES PROFESSIONNELLES")
    for i, exp in enumerate(CV["experiences"]):
        p = doc.add_paragraph(); set_spacing(p, before=5 if i else 1, after=0)
        p.paragraph_format.keep_with_next = True
        run(p, exp["dates"] + " | ", bold=True, color=ACCENT)
        run(p, exp["poste"] + " @ " + exp["entreprise"], bold=True)
        if exp["contrat"]:
            run(p, " " + exp["contrat"])
        if exp["perimetre"]:
            p = doc.add_paragraph(); set_spacing(p, after=2)
            run(p, exp["perimetre"], size=9, italic=True, color=GREY)
        for label, text, subs in exp["bullets"]:
            bullet(doc, label, text)
            for sub in subs:
                bullet(doc, "", sub, level=1)

    section(doc, "FORMATION")
    for dates, diplome, suite, ecole in CV["formation"]:
        p = doc.add_paragraph(); set_spacing(p, after=0 if ecole else 1)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(2.6))
        p.paragraph_format.left_indent = Cm(2.6)
        p.paragraph_format.first_line_indent = Cm(-2.6)
        run(p, dates, bold=True, color=ACCENT)
        run(p, "\t")
        run(p, diplome, bold=True)
        run(p, suite)
        if ecole:
            p = doc.add_paragraph(); set_spacing(p, after=2)
            p.paragraph_format.left_indent = Cm(2.6)
            run(p, ecole, italic=True, color=GREY)

    section(doc, "COMPÉTENCES")
    for label, text in CV["competences"]:
        p = doc.add_paragraph(); set_spacing(p, after=1, line=1.05)
        run(p, label, bold=True)
        run(p, text)

    section(doc, "LANGUES")
    p = doc.add_paragraph(); set_spacing(p, after=0)
    run(p, CV["langues"])

    section(doc, "CENTRES D'INTÉRÊT")
    p = doc.add_paragraph(); set_spacing(p, after=0)
    run(p, CV["interets"])

    out_dir = ROOT / "output"
    out_dir.mkdir(exist_ok=True)
    docx_path = out_dir / f"{out_stem}.docx"
    doc.save(docx_path)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
                   check=True, capture_output=True)
    return docx_path, out_dir / f"{out_stem}.pdf"


if __name__ == "__main__":
    stem = sys.argv[1] if len(sys.argv) > 1 else "CV_NicolasQiGUO_General"
    for path in build(stem):
        print(path)
